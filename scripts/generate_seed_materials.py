import openpyxl
import re

def clean_str(val):
    if val is None:
        return None
    s = str(val).strip()
    return s if s else None

def sql_escape(s):
    if s is None:
        return 'null'
    return "'" + str(s).replace("'", "''") + "'"

def main():
    wb = openpyxl.load_workbook('Ref files/RM-Finish Goods SKU.xlsx', data_only=True)
    out_lines = []
    out_lines.append("-- Seed real raw materials & coils from Ref files/RM-Finish Goods SKU.xlsx")
    out_lines.append("-- Run in Supabase SQL Editor after migrate-019.sql\n")

    # 1. Raw Material Coils
    ws_coils = wb['Raw Material Coils SKU']
    coil_values = []
    seen_codes = set()

    for r in range(3, ws_coils.max_row + 1):
        coil_code = clean_str(ws_coils.cell(r, 1).value)
        color_name = clean_str(ws_coils.cell(r, 2).value)
        size_str = clean_str(ws_coils.cell(r, 3).value)
        paint_type = clean_str(ws_coils.cell(r, 4).value)
        mat_desc = clean_str(ws_coils.cell(r, 5).value)
        uom = clean_str(ws_coils.cell(r, 6).value) or 'KG'

        if not mat_desc:
            continue
        if mat_desc in seen_codes:
            continue
        seen_codes.add(mat_desc)

        thickness = None
        width = None
        if size_str and '*' in size_str:
            parts = size_str.split('*')
            try:
                thickness = float(parts[0])
                width = float(parts[1])
            except ValueError:
                pass

        is_bot = (
            coil_code == 'BOTTOM COIL' or
            'bottom' in (color_name or '').lower() or
            'bottom' in mat_desc.lower() or
            'service' in mat_desc.lower()
        )
        kind = 'BOTTOM_COIL' if is_bot else 'TOP_COIL'
        name = f"{color_name or ''} {'Bottom ' if is_bot else ''}Coil ({size_str or ''})".strip()

        coil_values.append(
            f"({sql_escape(mat_desc)}, {sql_escape(name)}, 'RM', {sql_escape(paint_type)}, "
            f"{thickness if thickness is not None else 'null'}, {width if width is not None else 'null'}, "
            f"{sql_escape(uom)}, {sql_escape(kind)}, {sql_escape(coil_code if not is_bot else None)}, "
            f"{sql_escape(color_name)}, 0)"
        )

    # 2. Raw Material SKU (granules, adhesive, protective films)
    ws_rm = wb['Raw Material SKU']
    rm_values = []
    for r in range(3, ws_rm.max_row + 1):
        rm_name = clean_str(ws_rm.cell(r, 1).value)
        grade = clean_str(ws_rm.cell(r, 2).value)
        size_str = clean_str(ws_rm.cell(r, 3).value)
        brand = clean_str(ws_rm.cell(r, 4).value)
        mat_desc = clean_str(ws_rm.cell(r, 5).value)
        uom = clean_str(ws_rm.cell(r, 6).value) or 'KG'

        if not mat_desc or mat_desc in seen_codes:
            continue
        seen_codes.add(mat_desc)

        thickness = None
        width = None
        if size_str and '*' in size_str:
            parts = size_str.split('*')
            try:
                thickness = float(parts[0])
                width = float(parts[1])
            except ValueError:
                pass

        category = 'RM'
        kind = 'OTHER'
        desc_lower = mat_desc.lower()

        if 'adhesive' in desc_lower:
            kind = 'ADHESIVE'
            category = 'CONSUMABLE'
        elif 'protective' in desc_lower or 'masking' in desc_lower:
            kind = 'MASKING'
            category = 'CONSUMABLE'
        elif 'granule' in desc_lower or 'core' in desc_lower:
            kind = 'CORE'
            category = 'GRANULE' if 'granule' in desc_lower else 'RM'

        rm_values.append(
            f"({sql_escape(mat_desc)}, {sql_escape(mat_desc)}, {sql_escape(category)}, {sql_escape(grade)}, "
            f"{thickness if thickness is not None else 'null'}, {width if width is not None else 'null'}, "
            f"{sql_escape(uom)}, {sql_escape(kind)}, null, null, 0)"
        )

    all_values = coil_values + rm_values
    out_lines.append(f"-- Inserting {len(all_values)} materials ({len(coil_values)} coils + {len(rm_values)} other raw materials)")
    
    # Chunk into batches of 200
    chunk_size = 200
    for i in range(0, len(all_values), chunk_size):
        chunk = all_values[i:i+chunk_size]
        out_lines.append(
            "insert into materials (code, name, category, mat_type, thickness, width, uom, kind, color_code, color_name, cost_price)\n"
            "values\n" + ",\n".join(chunk) + "\n"
            "on conflict (code) do update set\n"
            "  kind = excluded.kind,\n"
            "  color_code = excluded.color_code,\n"
            "  color_name = excluded.color_name,\n"
            "  thickness = excluded.thickness,\n"
            "  width = excluded.width,\n"
            "  uom = excluded.uom;\n"
        )

    # Link standard materials to core_settings
    out_lines.append("""
-- Link standard core, adhesive, and masking materials to core_settings
update core_settings
   set core_material_id = (select id from materials where code = 'LDPE RECYCLED GRANULE---A' limit 1),
       adh_material_id = (select id from materials where code like 'ADHESIVE FILM--0.050*1280-NORMAL%' limit 1),
       mask_material_id = (select id from materials where code like 'PROTECTIVE FILM-ALUMAX-0.080*1260-LOW%' limit 1)
 where core_type = 'NFR';

update core_settings
   set core_material_id = (select id from materials where code like 'FR3 GRANULES (FRB1)%' limit 1),
       adh_material_id = (select id from materials where code like 'ADHESIVE FILM--0.050*1280-NORMAL%' limit 1),
       mask_material_id = (select id from materials where code like 'PROTECTIVE FILM-ALUMAX-0.080*1260-MEDIUM%' limit 1)
 where core_type = 'FRB1';

update core_settings
   set core_material_id = (select id from materials where code like 'FR2 GRANULE (FRA2 GRANULE)%' or code like 'FR2 CORE%' limit 1),
       adh_material_id = (select id from materials where code like 'ADHESIVE FILM --0.080*1280-FRA2%' limit 1),
       mask_material_id = (select id from materials where code like 'PROTECTIVE FILM-ALUMAX-0.080*1260-HIGH%' limit 1)
 where core_type = 'FRA2';
""")

    with open('tas-oms/supabase/seed_materials.sql', 'w', encoding='utf-8') as f:
        f.write('\n'.join(out_lines))
    print(f"Generated tas-oms/supabase/seed_materials.sql with {len(all_values)} materials.")

if __name__ == '__main__':
    main()
